"""Post-render fact checks and concurrent LLM/VLM review of actual documents."""
import base64
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pymupdf

from .io import read_json, write_json, file_hash
from .model_runtime import model_call, parallel_map, telemetry
from .facts import tunnel_mark


def collect_documents(out):
    documents=[{'id':'description','path':out/'路径方案描述_联调样章.docx','sections':['3.3.1','3.3.2']}]
    job_path=out/'agents/agent_run.json'
    if job_path.exists():
        job=read_json(job_path)
        for agent,info in job['results'].items():
            result=read_json(info['result_json'])
            for index,item in enumerate(info['documents']):
                if any(Path(item['docx']).resolve()==Path(d['path']).resolve() for d in documents):
                    continue
                documents.append({'id':f'{agent}_{index}','path':Path(item['docx']),
                                  'sections':result['agent']['sections']})
    assembly=out/'agents/assembly/result.json'
    if assembly.exists():
        value=read_json(assembly)
        for item in value['documents']:
            documents.append({'id':'assembly_'+item['scheme_id'],'path':Path(item['docx']),'sections':value['sections']})
    return documents


def reference_page(document_id,text):
    if document_id=='description':return 3
    if document_id.startswith('route_compliance_'):
        sid=list(evidence.get('compliance',{}))[int(document_id.rsplit('_',1)[1])]
        return {'schemes':[x for x in evidence['schemes'] if x['id']==sid],
                'compliance':{sid:evidence['compliance'][sid]},
                'document_scope':{'scheme_id':sid,'note':'本分册只描述此方案，不要求披露其他方案内容'},
                **({'presentation':evidence['presentation']} if 'presentation' in evidence else {})}
    if document_id.startswith('route_comparison'):

        return 5 if '综合' in text or '模型总成本' in text else 3
    if '3.5' in text or '协议办理' in text:return 23
    if any(s in text for s in ('3.4.10','3.4.11','林木','生态敏感')):return 22
    return 16


def image_url(pix):
    return 'data:image/png;base64,'+base64.b64encode(pix.tobytes('png')).decode()


def empty_body_pages(pdf):
    """Ignore the running header/footer: those alone do not make a report page."""
    empty=[]
    for index,page in enumerate(pdf):
        lower,upper=75,page.rect.height-65
        text=any(lower<(b[1]+b[3])/2<upper and b[4].strip() for b in page.get_text('blocks') if b[6]==0)
        images=any(lower<(i['bbox'][1]+i['bbox'][3])/2<upper for i in page.get_image_info())
        if not text and not images:empty.append(index+1)
    return empty


def check_review_consistency(config, result):
    """Reject self-contradictory review labels, preserving both model opinions."""
    blocking=[i for i,x in enumerate(result['issues']) if x['severity']=='blocking']
    if not blocking:return result,None
    decision,usage=model_call(dict(config,_purpose='review_consistency'),[
        {'role':'system','content':'你复核校对意见自身的逻辑一致性，不审批线路、不重新评判工程风险。仅当某条blocking意见的描述本身明确认定成稿准确、如实披露、不构成错误时，判为self_contradictory；若已承认事实/数值准确，仅建议重复提示、换表头、增加脚注或避免读者可能误解，判为suggestion_only（降为minor，保留建议）；如果指出真实错误数值、冒称合规或未披露实际冲突，或仍不确定，必须保留为coherent。不得因为缺工程资料就认定校对意见错误。按输入blocking索引逐一返回JSON：{"decisions":[{"index":0,"classification":"self_contradictory、suggestion_only或coherent","reason":"引述该条意见中自相矛盾的具体词句，或保留原因"}]}。'},
        {'role':'user','content':__import__('json').dumps({'review':result,'blocking_indices':blocking},ensure_ascii=False)}])
    rows=decision.get('decisions',[])
    if len(rows)!=len(blocking) or {x.get('index') for x in rows}!=set(blocking):
        raise ValueError('审阅一致性复核未逐项返回有效索引')
    if any(x.get('classification') not in ('self_contradictory','suggestion_only','coherent') or not x.get('reason') for x in rows):
        raise ValueError('审阅一致性复核格式无效')
    removed={x['index'] for x in rows if x['classification']=='self_contradictory'}
    import copy
    revised=copy.deepcopy(result);revised['issues']=[x for i,x in enumerate(result['issues']) if i not in removed]
    for old_index,issue in enumerate(result['issues']):
        if any(x['index']==old_index and x['classification']=='suggestion_only' for x in rows):
            next(x for x in revised['issues'] if x['description']==issue['description'])['severity']='minor'
    revised['verdict']='needs_review' if any(x['severity']=='blocking' for x in revised['issues']) else 'pass'
    return revised,{'original_review':result,'decisions':rows,'usage':usage,
                    'scope':'仅复核校对意见自身是否自相矛盾；未改变报告内容、几何冲突或工程状态'}


def document_evidence(evidence, document_id):
    # A description chapter is not a compliance decision; avoid asking the model
    # to treat unrelated JSON findings as statements in that document.
    if document_id == 'description':
        return {k:evidence[k] for k in ('schemes','presentation') if k in evidence}
    if document_id.startswith('route_compliance_'):
        sid=list(evidence.get('compliance',{}))[int(document_id.rsplit('_',1)[1])]
        return {'schemes':[x for x in evidence['schemes'] if x['id']==sid],
                'compliance':{sid:evidence['compliance'][sid]},
                'document_scope':{'scheme_id':sid,'note':'本分册只描述此方案，不要求披露其他方案内容'},
                **({'presentation':evidence['presentation']} if 'presentation' in evidence else {})}
    if document_id.startswith('route_comparison'):

        return {k:v for k,v in evidence.items() if k != 'compliance'}
    return evidence


def review_reports(config, out, retry_ids=None):
    out=Path(out);folder=out/'quality_review';folder.mkdir(parents=True,exist_ok=True)
    prior=read_json(out/'quality_review.json') if retry_ids is not None else None
    facts=read_json(out/'route_facts.json')
    comparisons=out/'agents/route_comparison/result.json'
    compliance=out/'agents/route_compliance/result.json'
    evidence={'schemes':[{'id':s['id'],'name':s['name'],'metrics':s['metrics'],
                          'counties':s['counties'],'villages':s.get('villages',[]),'admin_complete':s['admin_complete'],
                          'segments':s['segments'],'terrain':s.get('terrain',[]),'physical_parallel':s.get('physical_parallel',[]),
                          'shared_ends':s.get('shared_ends',{}),
                          'display_values':{'description_length_km':f"{s['metrics']['length_m']/1000:.1f}",
                                            'description_straight_km':f"{s['metrics']['straight_distance_m']/1000:.1f}",
                                            'description_tortuosity':f"{s['metrics']['tortuosity']:.2f}",
                                            'comparison_length_km':f"{s['metrics']['length_m']/1000:.3f}",
                                            'comparison_tortuosity':f"{s['metrics']['tortuosity']:.3f}"},
                          'events':[dict({k:e.get(k) for k in ('fact_id','label','category','start_m','end_m','relation','tunnel','source_uids')},tunnel_is_true=tunnel_mark(e.get('tunnel'))) for e in s['events']]}
                         for s in facts['schemes']]}
    if config.get('report_style')=='reference':
        for item in evidence['schemes']:item.pop('terrain',None)
        evidence['presentation']={'style':'reference','scope':'正文按参考报告介绍主要道路、电力线、铁路、河流和生态范围；普通乡道、小路、技术明细留在JSON/CSV，正文无需逐项复述。crossings中presentation_included=false的事件不计入正文主要跨越数；同图幅可涵盖多个邻近事件，不代表同一工程跨越档。未在正文重复辅助指标或每个事件不是事实错误。'}
    if comparisons.exists():
        result=read_json(comparisons)
        evidence['comparison']={'matrix':result['matrix'],'recommendation':result['recommendation'],'specialists':result.get('specialists',{}),'coverage':result.get('coverage',[])}
    if compliance.exists():
        result=read_json(compliance)
        evidence['compliance']={k:{'status':v['status'],'checks':v['checks'],
                                    'missing':v['missing'],'agreements':v['agreements'],
                                    'demolition':v['demolition'],'forestry':v['forestry'],'ecology_opinion':v.get('ecology_opinion'),
                                    'ecology':v.get('ecology',[]),'ecology_context':v.get('ecology_context',[]),'crossings':v['crossings']}
                                for k,v in result['schemes'].items()}
    modes=read_json(out/'narrative.json')
    checks=[{'check':'two_distinct_named_schemes','passed':len(facts['schemes'])==2 and {s['name'] for s in facts['schemes']}=={'北方案','南方案'}},
            {'check':'description_real_model','passed':all(v['mode']=='llm_constrained' for v in modes.values()) if config.get('llm',{}).get('enabled') else True}]
    if comparisons.exists():
        r=read_json(comparisons)
        checks.append({'check':'comparison_real_model','passed':r['model']['mode']=='llm_constrained' and all(v['mode']=='llm_constrained' for v in r['descriptions'].values()) if config.get('llm',{}).get('enabled') else True})
    if compliance.exists():
        r=read_json(compliance)
        checks.append({'check':'compliance_real_model','passed':all(v['model']['mode']=='llm_constrained' for v in r['schemes'].values()) if config.get('llm',{}).get('enabled') else True})
    documents=[];jobs=[]
    ref=pymupdf.open(config['reference_pdf'])
    reference_style={'page_size_pt':[ref[min(3,len(ref)-1)].rect.width,ref[min(3,len(ref)-1)].rect.height],
                     'body_font':'SimSun / 宋体','body_size_pt':12,'line_spacing_pt':24,
                     'source_pages':[p for p in [4,6,17,23,24] if p<=len(ref)] or [1],
                     'reference_page_count':len(ref),
                     'reference_scope':'当前参考PDF仅包含停电补充页；其余章节按已有模板核查' if len(ref)<4 else '更新后的完整参考章节PDF（含停电补充页）',
                     'comparison_scope':'页面尺寸、声明字体、图表布局；新路线内容/图源/分页允许与参考不同'}
    max_pages=int(config.get('quality_review',{}).get('max_pages',24))
    count=0
    for document in collect_documents(out):
        path=document['path'];doc=Document(path)
        headings='\n'.join(p.text for p in doc.paragraphs)
        with ZipFile(path) as archive:
            xml=archive.read('word/document.xml').decode()
        section_ok=all(section in headings for section in document['sections'])
        checks.append({'check':document['id']+'_sections_and_fonts','passed':section_ok and '宋体' in xml and 'Times New Roman' in xml})
        pdf_path=path.with_suffix('.pdf');pdf=pymupdf.open(pdf_path)
        actual_text='\n'.join(page.get_text() for page in pdf)
        record={'id':document['id'],'docx':str(path),'pdf':str(pdf_path),'pages':len(pdf),
                'docx_sha256':file_hash(path),'pdf_sha256':file_hash(pdf_path),'tables':len(doc.tables)}
        documents.append(record)
        overflow=[]
        for n,page in enumerate(pdf):
            for block in page.get_text('dict')['blocks']:
                if block.get('type')!=0:continue
                for line in block['lines']:
                    for span in line['spans']:
                        x0,y0,x1,y1=span['bbox']
                        if x0 < -1 or y0 < -1 or x1>page.rect.width+1 or y1>page.rect.height+1:
                            overflow.append({'page':n+1,'bbox':[x0,y0,x1,y1]})
        checks.append({'check':document['id']+'_page_bounds','passed':not overflow,'details':overflow})
        empty=empty_body_pages(pdf)
        checks.append({'check':document['id']+'_no_empty_body_pages','passed':not empty,'details':empty})
        checks.append({'check':document['id']+'_reference_a4','passed':all(abs(page.rect.width-reference_style['page_size_pt'][0])<2 and abs(page.rect.height-reference_style['page_size_pt'][1])<2 for page in pdf)})
        if config.get('llm',{}).get('enabled'):
            jobs.append({'id':document['id']+'_text','kind':'text','config':dict(config['llm'],_purpose='review_text'),
                'messages':[{'role':'system','content':'你是输电线路报告事实校对员。仅核对提供的证据和成稿，不补造业务事实。正文长度保留0.1km、曲折系数2位小数，比较表分别3位小数；精确显示值见display_values，正常四舍五入不是错误。分段长度及平行长度也保留0.1km。缺照片、复函、拆迁调查等明确写为待补是正确行为。返回纯JSON：{"verdict":"pass或needs_review","issues":[{"severity":"blocking或minor","description":"具体矛盾及位置"}],"summary":"简述"}。自然语言内部引用只能用中文「」或不加引号，不能有未转义的英文双引号。只判断document_text中的实际陈述，evidence中的JSON键值不是成稿原文。校对报告是否如实记录风险，不负责批准路线；报告包含真实规则冲突不是报告写错。若正文已指出具体对象、里程、角度和未满足规则，即为已披露；不要求在每一张表重复披露，也不能因方案有冲突而要求从比较表删掉它。表中列出对象不等于声称其合规。每条blocking必须说明成稿哪句错误陈述与哪项证据相反，不能把已如实描述的风险自身列为blocking。已明确披露为需复核的规则冲突不等于成稿错误，不得要求擅自改成已合规或正式否决；未披露而冒称无冲突才是错误。只有证据冲突、错误数值、冒称已批准才是blocking。无实质问题时issues必须为空数组，正确、一致、允许的取整不得列为问题。'},
                            {'role':'user','content':__import__('json').dumps({'evidence':document_evidence(evidence,document['id']),'document_text':actual_text},ensure_ascii=False)}]})
        if config.get('quality_review',{}).get('vision') and config.get('vlm',{}).get('enabled'):
            for n,page in enumerate(pdf):
                if count>=max_pages:break
                count+=1
                jobs.append({'id':document['id']+f'_page_{n+1:02d}','kind':'vision','config':dict(config['vlm'],_purpose='review_page'),
                    'messages':[{'role':'system','content':'你只检查一张实际Word预览页面的可读性，不审核工程内容完整性。检查明显文字重叠、裁切、无法阅读、图题与图片挤压；只报告眼睛能看到的问题。GIS矢量图、留白、待核实文字、与别的报告不同的地名和页码都是允许的，不得要求卫星图/现场照片或复制其他报告的内容。返回JSON：{"verdict":"pass或needs_review","issues":[{"severity":"blocking或minor","description":"具体位置和可见问题"}],"summary":"简述"}。只有明显影响阅读的可见缺陷为blocking。裁切必须看到字形被页面边界或遮挡截断，不能根据语意猜测图题还应多几个字。同一章节分别列出南北方案的同名表格是允许的，重复内容不属于视觉排版错误。'},
                                {'role':'user','content':[{'type':'text','text':f"请只检查当前文档{document['id']}第{n+1}页是否排版可读。它是新路线的核查稿，目标为A4、宋体12pt正文、规范图题表头；参考PDF样式参数已由程序单独对照。不要评判缺照片或待补业务资料。只输出上述JSON对象。"},
                                  {'type':'image_url','image_url':{'url':image_url(page.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25)))}}]}]})
        pdf.close()
    ref.close()
    def execute(job):
        try:
            result,usage=model_call(job['config'],job['messages'])
            if result.get('verdict') not in ('pass','needs_review') or not isinstance(result.get('issues'),list):
                raise ValueError('审阅响应格式无效')
            for issue in result['issues']:
                if issue.get('severity') not in ('blocking','minor') or not isinstance(issue.get('description'),str):
                    raise ValueError('审阅问题格式无效')
            consistency=None
            if job['kind']=='text':result,consistency=check_review_consistency(job['config'],result)
            value={'id':job['id'],'kind':job['kind'],'status':'completed','result':result,'usage':usage}
            if consistency:value['consistency_review']=consistency
        except Exception as exc:
            value={'id':job['id'],'kind':job['kind'],'status':'failed','error_type':type(exc).__name__}
        write_json(folder/(job['id']+'.json'),value)
        print('审阅完成',job['id'],value['status'],flush=True)
        return value
    if retry_ids is not None:
        # Explicit targeted retry only; never reuse reviews after any document changes.
        if prior['documents'] != documents:
            raise ValueError('成稿哈希/页数已改变，不能仅重试部分审阅')
        wanted=set(retry_ids);known={j['id'] for j in jobs}
        if not wanted <= known:raise ValueError('重试审阅ID不存在')
        saved={r['id']:r for r in prior['reviews']}
        if any(j['id'] not in saved for j in jobs if j['id'] not in wanted):
            raise ValueError('缺少可复用的原审阅记录')
        history=folder/'history';history.mkdir(exist_ok=True)
        import time
        stamp=str(time.time_ns())
        write_json(history/('quality_review_'+stamp+'.json'),prior)
        fresh=parallel_map(execute,[j for j in jobs if j['id'] in wanted],config.get('llm',{}).get('max_concurrency',15))
        saved.update({r['id']:r for r in fresh});results=[saved[j['id']] for j in jobs]
    else:
        results=parallel_map(execute,jobs,config.get('llm',{}).get('max_concurrency',15))
    blocked=[r['id'] for r in results if r['status']=='completed' and any(i['severity']=='blocking' for i in r['result']['issues'])]
    failed=[r['id'] for r in results if r['status']!='completed']
    missing_pages=max(0,sum(d['pages'] for d in documents)-count)
    if config.get('quality_review',{}).get('vision') and config.get('vlm',{}).get('enabled'):
        checks.append({'check':'all_document_pages_reviewed','passed':missing_pages==0})
    result={'status':'passed' if all(c['passed'] for c in checks) and not blocked and not failed else 'needs_review',
            'checks':checks,'documents':documents,'reviews':results,'blocking_reviews':blocked,'failed_reviews':failed,
            'vision_pages_reviewed':count,'vision_pages_not_reviewed':missing_pages,
            'scope':'报告事实与可读性检查，不代替缺失的工程调查和正式审批',
            'reference_style':reference_style,
            'reference_pdf':config['reference_pdf'],'reference_pdf_sha256':file_hash(config['reference_pdf'])}
    write_json(out/'quality_review.json',result)
    model_log=out/('model_calls_review_retry.json' if retry_ids is not None else 'model_calls.json')
    if model_log.exists():
        import time
        history=out/'model_call_history';history.mkdir(exist_ok=True)
        write_json(history/(model_log.stem+'_'+str(time.time_ns())+'.json'),read_json(model_log))
    write_json(model_log,telemetry())
    return result
