"""Repeatable chapter-by-chapter delivery audit; --live probes only this session."""
import argparse
from collections import Counter
from datetime import datetime
from pathlib import Path
import sys
import uuid

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
import pymupdf
from road_description.io import read_json,write_json,file_hash,public_data, config_root
from road_description.planner import PlannerClient,PlannerError,select_scheme
from road_description.quality_review import collect_documents


def probe_interfaces(config, root, out):
    client=PlannerClient(config['planner_url'])
    results=[]
    def probe(name,action,valid):
        try:
            data=action(); ok=valid(data)
            results.append({'interface':name,'passed':bool(ok),'response':public_data(data)})
        except Exception as exc:
            results.append({'interface':name,'passed':False,'error_type':type(exc).__name__,
                            'diagnostic':getattr(client,'last_failure',None)})
    probe('GET /api/health',lambda:client.call('GET','/api/health'),lambda x:x.get('status') in ('ok','healthy'))
    # A full map is intentionally kept in memory only; archive key metadata and counts.
    def map_summary():
        _,data=client.initialize()
        return {k:data.get(k) for k in ('preprocess_mode','geo_meta','manifest_path','manifest_mtime','features_lazy')}|{
                'feature_count':len(data.get('features',[])),'session_cookie_present':bool(client.session.cookies)}
    probe('GET /api/map',map_summary,lambda x:x['preprocess_mode'] and x['geo_meta'] and x['session_cookie_present'])
    probe('GET /api/map/meta',client.map_meta,lambda x:bool(x.get('map_cache_key')))
    start,end=config['start'],config['end']
    bounds=[min(start[0],end[0])-100,min(start[1],end[1])-100,max(start[0],end[0])+100,max(start[1],end[1])+100]
    def features():
        data=client.map_features(bounds,max_features=5,tol=0)
        return {k:data.get(k) for k in ('returned','total_in_view','capped','tol')}|{
            'sample_feature_keys':[sorted(f.keys()) for f in data.get('features',[])],
            'geometry_present':all(bool(f.get('coords') or f.get('geometry') or f.get('points')) for f in data.get('features',[]))}
    probe('GET /api/map/features',features,lambda x:isinstance(x['returned'],int) and x['returned']>0 and x['geometry_present'])
    for name,point in [('start',start),('end',end)]+[(s['id'],s['waypoint']) for s in config['schemes'] if s.get('waypoint')]:
        probe('GET /api/validate_constraint_point '+name,lambda p=point:client.validate(p),lambda x:x.get('valid') is True)
    token='rd_audit_absent_'+uuid.uuid4().hex
    probe('GET /api/plan/progress (本审计会话)',lambda:client.progress(token),lambda x:x.get('status') in ('idle','stale'))
    probe('POST /api/plan/cancel (不存在的专属ID)',lambda:client.cancel(token),lambda x:x.get('cancelled') is False)
    def invalid():
        try:client.call('GET','/api/validate_constraint_point',params={'x':'invalid','y':0})
        except PlannerError:return client.last_failure or {}
        return {}
    probe('非法控制点返回受控400',invalid,lambda x:x.get('status_code')==400)
    probe('GET /api/projects (扩展入口可达)',lambda:{'project_count':len(client.call('GET','/api/projects')['projects'])},lambda x:x['project_count']>=0)
    result={'checked_at':datetime.now().isoformat(),'checks':results,'passed':all(r['passed'] for r in results),
            'scope':'报告主链接口实测；取消仅对全新会话的不存在任务ID。工程项目异步任务提交/执行不属于本次报告调用链，未新建工程或冒称实跑。'}
    write_json(out/'interface_audit.json',result)
    return result


def audit(config_path,live=False):
    config_path=Path(config_path).resolve();root=config_root(config_path);c=read_json(config_path)
    out=(root/c['output_dir']).resolve()
    interfaces=probe_interfaces(c,root,out) if live else read_json(out/'interface_audit.json')
    facts=read_json(out/'route_facts.json');narratives=read_json(out/'narrative.json')
    comp=read_json(out/'agents/route_comparison/result.json');compliance=read_json(out/'agents/route_compliance/result.json')
    job=read_json(out/'agents/agent_run.json');review=read_json(out/'quality_review.json')
    checks=[];rows=[];documents=[]
    def check(label,ok,details=None):checks.append({'check':label,'passed':bool(ok),'details':details})
    def row(agent,chapter,requirement,status,evidence,gap=''):
        rows.append(dict(agent=agent,chapter=chapter,requirement=requirement,status=status,evidence=evidence,gap=gap))
    check('三个插件均已执行',set(job['results'])=={'route_description','route_comparison','route_compliance'})
    for item in collect_documents(out):
        path=Path(item['path']);doc=Document(path);pdf=pymupdf.open(path.with_suffix('.pdf'))
        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(cell.text for t in doc.tables for r in t.rows for cell in r.cells)
        record={'id':item['id'],'docx':str(path),'pdf':str(path.with_suffix('.pdf')),'pages':len(pdf),
                'tables':len(doc.tables),'images':len(doc.inline_shapes),'text_characters':len(text),
                'docx_sha256':file_hash(path),'pdf_sha256':file_hash(path.with_suffix('.pdf'))}
        documents.append(record)
        check(item['id']+'章节标题',all(section in text for section in item['sections']))
        check(item['id']+'已审阅当前版本',any(r['docx_sha256']==record['docx_sha256'] and r['pdf_sha256']==record['pdf_sha256'] for r in review['documents']))
        pdf.close()
    check('质量审阅覆盖全部页面',review['status']=='passed' and review['vision_pages_not_reviewed']==0)
    check('当前主链接口检查',interfaces['passed'])
    for index,s in enumerate(facts['schemes'],1):
        sid=s['id'];chapter=f'3.3.{index}';name=s['name'];n=narratives[sid]
        ids={v for p in n['paragraphs'] for v in p['fact_ids']}
        check(name+'所有地物事件均入正文',{e['fact_id'] for e in s['events']}.issubset(ids))
        check(name+'所有路段均入正文',{e['fact_id'] for e in s['segments']}.issubset(ids))
        check(name+'长度与API一致',abs(s['metrics']['length_m']-s['metrics']['api_length_m'])<.11)
        response=read_json(out/sid/'response.json');selected=select_scheme(response)
        check(name+'真实完整规划及来源',bool(response.get('plan_run_id')) and selected['feasible'] is True and bool(response.get('report_selection',{}).get('source_sha256')))
        source=response.get('report_selection',{})
        original=Path(source.get('source_response',''))
        check(name+'原始响应哈希未改变',original.is_file() and file_hash(original)==source.get('source_sha256'))
        # Response provenance remains in the original planning_attempts directory.
        row('路径描述',chapter,'起终点、逐段走向/长度、总长/曲折系数','已实算',f"{name}：{len(s['segments'])}段，{s['metrics']['length_m']/1000:.3f}km；response.json与route_facts.json")
        row('路径描述',chapter,'经过地物及重复交叉的位置','已实算',f"{len(s['events'])}个事件均入正文，保留source_uids和里程")
        row('路径描述',chapter,'平行区段、共同段、规则与避让依据','已实算',f"平行记录{len(s.get('physical_parallel',[]))}条；邻近禁行记录{len(s['constraints'])}条；shared_ends已计算",'仅描述地图中实际存在的关系，不能复制原PDF的20km等数字')
        row('路径描述',chapter,'正式地名及完整行政区县统计','部分具备','有专名则使用原属性；其他用类别及里程定位','完整行政区划、未命名地物专名及正式站名待甲方确认' if not s['admin_complete'] else '')
        row('路径描述',chapter,'总览与沿线图片','部分具备',f"已生成{len(s['figure_ids'])}张沿线GIS图及双方案总览",'现场照片/获授权航拍影像及拍摄点资料待补；GIS图不冒充实景')
        result=compliance['schemes'][sid];events=result['crossings'];figures=result['figures']
        maps=[f for f in figures if f['kind'] in ('vector_map','satellite_map')]
        profiles=[f for f in figures if f['kind']=='dem_terrain_profile']
        check(name+'跨越均有平面图',{e['event_id'] for e in events}=={f['event_id'] for f in maps})
        check(name+'跨越均有真实DEM剖面',{e['event_id'] for e in events}=={f['event_id'] for f in profiles})
        check(name+'跨越图保留原始地物ID',all(f.get('source_uids') for f in maps))
        # Every compliance crossing must also appear in the descriptive facts.
        check(name+'描述与合规跨越对应',all(any(set(e['source_uids'])&set(d['source_uids']) and abs(e['start_m']-d['start_m'])<1 for d in s['events']) for e in events))
        row('合规分析','3.4.8.1',name+'主要交叉跨越一览及逐项说明','已实算',f"{len(events)}个事件；{len(maps)}张平面图；分类/名称、里程、几何角及规则阈值")
        row('合规分析','3.4.8.1',name+'地形与工程跨越断面','部分具备',f"{len(profiles)}张DEM地形剖面，保留采样点与DEM哈希",'需正式塔位塔型、耐张段长、跨越档距、被跨物标高、导地线弧垂、净空及接头要求；当前不是设计断面')
        for chapter,key,title,missing in [('3.4.8.2','outage','停电方案','调度确认、停电对象/窗口、运行方式、风险与施工安排'),
                ('3.4.9','demolition','拆迁分类调查表','房屋/棚、坟墓、通信塔的实测数量及调查确认范围'),
                ('3.4.10','forestry','林木砍伐情况及表格','林区长度、树种树高、砍伐面积、高跨净空及通道设计')]:
            status='已有确认资料' if result[key]['status']=='confirmed_input' else '待外部资料'
            row('合规分析',chapter,name+title,status,'正文与对应表格已保留；result.json记录状态',missing if status=='待外部资料' else '')
        row('合规分析','3.4.11',name+'生态敏感点相对位置、影响及处理措施','部分具备',f"原始生态边界筛查{len(result['ecology'])}条记录，含距离/相交长度及系统规则",'施工占地、完整保护范围、逐塔影响和正式专题措施待补；零命中仅限当前图层/范围')
        row('合规分析','3.5',name+'行政区域/部门/办理状态/条件意见','已有确认资料' if result['agreements']['status']=='recorded_closed' else '待外部资料','协议汇总表及来源校验接口已具备','应办部门清单、来往函件与附条件意见落实情况待补')
    matrix=comp['matrix']['rows'];labels=[r[0] for r in matrix]
    expected=['线路长度','海拔分布','曲折系数','地形划分','沿线矿产','重要交叉跨越','房屋拆迁','穿越生态敏感','穿越生态保护红线','交通条件','军事设施','有关单位意见']
    check('比选覆盖参考表全部类别',all(any(k in label for label in labels) for k in expected))
    row('方案比选','3.3','影响因素、两走廊背景与总览','已实算','依据原始地物筛选两路线周边1000m影响因素，保留全清单及source_uids','矿区等专项层缺失不能按无矿解释')
    row('方案比选','3.3.1—3.3.2','两方案描述及图片','已实算','复用同一事实提取、文字组织和GIS绘图实现，分别有完整正文/图组')
    row('方案比选','3.3.3','综合比较表各指标','部分具备',f'{len(matrix)}个比较行；覆盖原PDF的长度、高程、曲折、地形、矿产、跨越、拆迁、生态、交通、军事及单位意见','地形分类比例、矿权专项、交通、军事、拆迁调查及正式复函待补')
    row('方案比选','3.3.3','指标差异与综合推荐','部分具备','已给出长度、模型成本、跨越/红线差异；recommendation保留证据范围','完整专题及专业审查未齐，正式综合推荐待定；不自动照抄北方案优选')
    row('三个Agent','共同排版','页面、字体、行距、图题表题与可编辑Word','部分具备','按参考设置A4、宋体/Times New Roman、正文12pt/24pt行距、页眉页码、双列沿线图及可编辑表格；程序另检查空白页、越界和图文排版',
        '本机PDF预览缺原字体时会使用替代字体；正式分页需在安装指定字体的Word/WPS中确认。现场照片和工程设计断面仍需实有资料。')
    ref=(root/c['reference_pdf']).resolve();pdf=pymupdf.open(ref)
    ref_pages=[3,4,5,6,7,8,17,18,19,20,21,22,23,24,25]
    (out/'参考PDF对应章节原文.txt').write_text('\n\n'.join(f'PDF文件页{p}（印刷页{p+13}）\n'+pdf[p-1].get_text() for p in ref_pages),encoding='utf-8');pdf.close()
    model=read_json(out/'model_calls.json')
    check('模型调用均成功且并发不超15',bool(model['calls']) and all(x['status']=='completed' for x in model['calls']) and model['peak_active']<=15)
    model_summary={'logical_calls':len(model['calls']),'http_attempts':sum(x['attempts'] for x in model['calls']),
                   'peak_active':model['peak_active'],'models':sorted({x.get('model_reported','unknown') for x in model['calls']})}
    result={'checked_at':datetime.now().isoformat(),'technical_status':'passed' if all(x['passed'] for x in checks) else 'needs_fix',
            'business_status':'draft_pending_external_evidence','reference_pdf':str(ref),'reference_sha256':file_hash(ref),
            'reference_pages':ref_pages,'agents':list(job['results']),'documents':documents,'checks':checks,'requirements':rows,
            'map_provenance':facts['provenance'],'interfaces':str(out/'interface_audit.json'),'model_summary':model_summary,
            'statement':'章节齐全与数据核验通过，不等于外部工程资料齐全。当前为约8km联调区段，非参考PDF正式全线。'}
    write_json(out/'chapter_audit.json',result)
    lines=['三个Agent逐项核验与资料缺口',f"核验时间：{result['checked_at']}",f"技术核验：{result['technical_status']}；业务状态：资料待补的核查稿",'',
        '一、交付文件（3个Agent，合规按南北分别出稿，共4份Word）']
    lines += [f"{d['id']}：{d['docx']}；{d['pages']}页，{d['tables']}表，{d['images']}图" for d in documents]
    lines += ['','二、逐项对应参考PDF（文件页3—8对应3.3；17—23对应3.4.8—3.4.11；23—25对应3.5）']
    for i,r in enumerate(rows,1):
        lines += [f"{i}. [{r['agent']} / {r['chapter']}] {r['requirement']} —— {r['status']}",'   依据：'+r['evidence']]
        if r['gap']:lines.append('   待补/边界：'+r['gap'])
    lines += ['','三、程序逐项核验']+[('通过' if x['passed'] else '失败')+'：'+x['check'] for x in checks]
    lines += ['','四、接口实测']+[('通过' if x['passed'] else '失败')+'：'+x['interface'] for x in interfaces['checks']]
    lines += [interfaces['scope'],'POST /api/plan：本轮复用两次已归档真实规划结果，校验完整终检与长度；未重复宣称本轮重新求解。',
        f"LLM/VLM：{model_summary['logical_calls']}次逻辑调用，{model_summary['http_attempts']}次HTTP请求，并发峰值{model_summary['peak_active']}；模型{','.join(model_summary['models'])}。正文组织、全文事实审阅与全部{sum(d['pages'] for d in documents)}页视觉审阅由本轮真实调用完成。",
        '','五、地图来源及适用范围','GIS原始地物、运行规则、禁行/禁立塔图层、坐标转换及DEM来源见chapter_audit.json / 各Agent result.json的哈希记录。',
        'HTML地图地名主要来自影像注记；现有属性未形成完整行政区划/地名对照表。不能把类别名称充作正式村镇名。',
        '规划接口可行不等于全部合规。当前统计塔模型未定位全部直线塔；起点终端塔需确认站址依据；规划响应DEM未加载，本报告DEM采样仅作独立分析。',
        '','六、需要甲方统一补充/确认的问题',
        '1. 正式项目起终点站名、坐标、线路电压、全线范围，以及南北走廊控制要求；当前8km试验不能代表原PDF44/47km全线。',
        '2. 路线照片、授权航拍/影像及位置/朝向/时间，是否接受现阶段GIS图作为过程稿插图？',
        '3. 杆塔逐塔坐标/塔型、耐张段、跨越档距、导地线参数/弧垂、被跨物标高、净空与接头要求，以及正式断面图。',
        '4. 停电方案/调度复函及其正式报告章节引用；不能直接套用原PDF“详见2.8节”。',
        '5. 拆迁、林业调查范围与分类数量，树种树高、砍伐量、通道及高跨设计，需绑定对应南北路线。',
        '6. 完整行政区划/正式地名、矿权、地形分类、军事设施及施工交通专项数据的权威来源/版本。',
        '7. 生态保护范围/专题成果、塔基和施工临时占地，及正式处理措施。',
        '8. 应办协议清单、部门复函、条件意见落实情况，以及最终推荐方案的专业审查口径。',
        '','复跑：taskset -c 16-19 .venv/bin/python -m road_description --config configs/config.chapter-audit.json',
        '核验：taskset -c 16-19 .venv/bin/python tools/audit_delivery.py --config configs/config.chapter-audit.json --live',
        '单独Agent：.venv/bin/python -m report_agents --config configs/config.chapter-audit.json --agents route_description（或route_comparison、route_compliance）',
        '单独执行会更新该稿；完整交付需要复跑整链审阅及核验，使文件哈希与审阅记录保持一致。']
    report='\n'.join(lines)+'\n'
    (out/'三Agent逐项核验与资料缺口.txt').write_text(report,encoding='utf-8')
    (root/'三Agent逐项核验与资料缺口_20261025.txt').write_text(report,encoding='utf-8')
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--config',default='configs/config.chapter-audit.json');parser.add_argument('--live',action='store_true')
    args=parser.parse_args();result=audit(args.config,args.live)
    print(result['technical_status'],len(result['checks']),'checks;',len(result['requirements']),'requirements')
    for check in result['checks']:
        if not check['passed']:print('FAILED:',check['check'])
    sys.exit(0 if result['technical_status']=='passed' else 2)
