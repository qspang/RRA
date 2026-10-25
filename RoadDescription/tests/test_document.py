from pathlib import Path
from zipfile import ZipFile
from docx import Document
from road_description.document import create_template,paragraph


def test_template_has_editable_text_and_reference_font_contract(tmp_path):
    path=tmp_path/'template.docx'
    create_template(path,{'project_title':'测试','page_start':17})
    doc=Document(path);paragraph(doc,'3.3.1 北方案路径',heading=True)
    paragraph(doc,'向西行进约2.5km。');doc.save(path)
    loaded=Document(path)
    assert any('2.5km' in p.text for p in loaded.paragraphs)
    with ZipFile(path) as archive:
        xml=archive.read('word/document.xml').decode()
        assert '宋体' in xml and 'Times New Roman' in xml
        assert 'w:start="17"' in xml
        assert 'w:line="480"' in xml


def test_renderer_does_not_claim_stale_pdf_as_new_success(tmp_path):
    from road_description.document import render_pdf
    docx=tmp_path/'report.docx';Document().save(docx)
    stale=tmp_path/'report.pdf';stale.write_bytes(b'old pdf')
    binary=tmp_path/'empty-renderer.sh'
    binary.write_text('#!/bin/sh\nexit 0\n');binary.chmod(0o755)
    result=render_pdf(docx,str(binary))
    assert result['status']=='failed'
    assert stale.read_bytes()==b'old pdf'
