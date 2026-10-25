"""Verify positive crossings in real planner JSON, Word content and mapped figures."""
import argparse,json,sys
from pathlib import Path
from collections import Counter
from docx import Document
from shapely.geometry import LineString,Point
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from road_description.io import read_json,write_json,file_hash
from report_agents.common.context import load_context
from report_agents.common.gis import is_ecology

def main():
    p=argparse.ArgumentParser();p.add_argument('--config',required=True);p.add_argument('--require_outside_endpoints',action='store_true');a=p.parse_args()
    ctx=load_context(a.config);out=ctx.source;check=[];summary=[]
    def yes(name,value):check.append({'check':name,'passed':bool(value)})
    result=read_json(ctx.output/'route_compliance/result.json')
    comparison=read_json(ctx.output/'route_comparison/result.json')
    cd=Document(comparison['documents'][0]['docx'])
    comparison_text='\n'.join(p.text for p in cd.paragraphs)+'\n'+'\n'.join(c.text for t in cd.tables for r in t.rows for c in r.cells)
    for fact in ctx.facts['schemes']:
        sid=fact['id'];r=result['schemes'][sid];line=LineString(fact['path']);counts=Counter(x['category'] for x in r['crossings'])
        positive=[e for e in r['ecology'] if e['inside_length_m']>.05]
        for word in ['高速','铁路','河流','输电']:
            yes(sid+'真实相交：'+word,any(word in k and v>0 for k,v in counts.items()))
        yes(sid+'真实生态正例',bool(positive))
        if a.require_outside_endpoints:
            yes(sid+'起终点均在生态范围外',all(not any(is_ecology(f) for f in ctx.catalog.nearby(Point(p))) for p in (line.coords[0],line.coords[-1])))
            yes(sid+'中段进入及离开生态范围',any(i['start_m']>1 and i['end_m']<line.length-1 for e in positive for i in e['intervals']))
        yes(sid+'生态图包含实际正例',all(any(f.get('ecology_source_uid')==e['source_uid'] for f in r['figures']) for e in positive))
        yes(sid+'每一交叉事件均配图',all(any(f.get('event_id')==e['event_id'] for f in r['figures']) for e in r['crossings']))
        document=Document(r['document']['docx'])
        text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for row in t.rows for c in row.cells)
        for word in ['高速','漳龙','生态','停电']:yes(sid+'实际Word含'+word,word in text)
        yes(sid+'实际Word含生态长度',all(f"{e['inside_length_m']:.1f}m" in text for e in positive))
        tunnels=[e for e in r['crossings'] if e.get('tunnel_status')=='tunnel']
        if tunnels:yes(sid+'隧道区别进入正文','不直接等同露天道路' in text)
        m=comparison['measurements'][sid]
        yes(sid+'风冰区真实覆盖',m['weather_zones']['covered_m']>0)
        yes(sid+'地块坐落单位村名已使用',bool(fact['villages']) and all(v['name'] in comparison_text for v in fact['villages']))
        yes(sid+'实际Word含风冰与高程起伏','风冰区资料' in comparison_text and '累计上升' in comparison_text and '覆冰' in comparison_text)
        conflicts=[e for e in r['crossings'] if e.get('angle_status')=='conflict']
        yes(sid+'实际冲突明确披露位置与数值',all(e['event_id'] in text and f"{e['angle_deg']:.2f}°" in text and '尚未满足当前规则' in text for e in conflicts))
        summary.append({'scheme':fact['name'],'length_km':line.length/1000,'crossings':dict(counts),
            'ecology':[{k:e[k] for k in ['label','inside_length_m','intervals','located_tower_count','tower_coverage_complete']} for e in positive],
            'villages':fact['villages'],'weather_zones':m['weather_zones'],'tunnel_events':len(tunnels),
            'report_status':r['status']})
    yes('交付检查通过',read_json(out/'delivery_manifest.json')['status']=='passed')
    yes('全文及逐页检查通过',read_json(out/'quality_review.json')['status']=='passed')
    assembly=next(d for d in read_json(ctx.output/'assembly/result.json')['documents'] if d['scheme_id']=='all')
    value={'status':'passed' if all(x['passed'] for x in check) else 'failed','checks':check,'schemes':summary,
           'full_report':assembly,'pdf_sha256':file_hash(assembly['pdf']['path']),
           'scope':'真实规划正例与输出一致性；几何事件不直接作为归并后的工程跨越处数，未替代设计/调查/审批'}
    write_json(out.parent/'positive_acceptance.json',value)
    print(json.dumps({'status':value['status'],'checks':len(check),'failed':[x for x in check if not x['passed']],
        'pages':assembly['pdf']['pages'],'schemes':summary},ensure_ascii=False,indent=2))
    if value['status']!='passed':raise SystemExit(1)

if __name__=='__main__':main()
