"""Editable DOCX using measured reference-PDF typography and two-column figures."""
import os
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image


def fonts(run,size=12,bold=False):
    from docx.shared import RGBColor
    run.font.color.rgb=RGBColor(0,0,0)
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold
    f=run._element.get_or_add_rPr().get_or_add_rFonts()
    f.set(qn('w:eastAsia'),'宋体');f.set(qn('w:ascii'),'Times New Roman');f.set(qn('w:hAnsi'),'Times New Roman')


def paragraph(document,text='',caption=False,heading=False):
    p=document.add_paragraph()
    fmt=p.paragraph_format
    fmt.line_spacing=Pt(24);fmt.space_after=Pt(0);fmt.space_before=Pt(0)
    fmt.first_line_indent=Pt(0 if caption else 24)
    fmt.widow_control=True
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if caption else WD_ALIGN_PARAGRAPH.JUSTIFY
    if heading:fmt.keep_with_next=True;fmt.space_before=Pt(6)
    fonts(p.add_run(text))
    return p


def extract_logo(reference,target):
    try:
        import pymupdf as fitz
        pdf=fitz.open(reference)
        for item in pdf[min(3,len(pdf)-1)].get_images(full=True):
            if item[2]>1000 and item[2]/item[3]>10:
                data=pdf.extract_image(item[0])
                target=Path(target).with_suffix('.'+data['ext']);target.parent.mkdir(parents=True,exist_ok=True)
                target.write_bytes(data['image']);pdf.close();return target
        pdf.close()
    except (ImportError,OSError,RuntimeError):
        pass
    return None


def create_template(path,config,logo=None):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.6);sec.bottom_margin=Cm(2.3)
    sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    sec.header_distance=Cm(1.35);sec.footer_distance=Cm(1.5)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.line_spacing=Pt(24)
    header=sec.header.paragraphs[0]
    header.paragraph_format.space_after=Pt(2)
    if logo:
        header.add_run().add_picture(str(logo),width=Cm(8.1))
    else:
        fonts(header.add_run('输电线路路径描述'),9)
    fonts(header.add_run('  路径方案报告 · 联调样章'),9)
    borders=OxmlElement('w:pBdr');border=OxmlElement('w:bottom')
    for k,v in [('val','single'),('sz','4'),('color','000000'),('space','4')]:border.set(qn('w:'+k),v)
    borders.append(border);header._p.get_or_add_pPr().append(borders)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fonts(footer.add_run('- '),9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    fonts(footer.add_run(' -'),9)
    page=OxmlElement('w:pgNumType');page.set(qn('w:start'),str(config.get('page_start',17)));sec._sectPr.append(page)
    doc.core_properties.title=config['project_title']
    doc.core_properties.subject='依据规划结果与GIS地物生成的路径描述联调样章'
    doc.core_properties.author='RoadDescription'
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True);doc.save(path)


def picture(p,path,width_cm):
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width_cm))


def figure_grid(doc,images,caption):
    table=doc.add_table(rows=0,cols=2);table.autofit=False
    for column in table.columns:column.width=Cm(8)
    for offset in range(0,len(images),2):
        row=table.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for j in range(2):
            if offset+j>=len(images):break
            item=images[offset+j];cell=row.cells[j]
            picture(cell.paragraphs[0],item['path'],7.6)
            p=cell.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing=Pt(18);p.paragraph_format.space_after=Pt(6)
            fonts(p.add_run(item['title']),11)
    paragraph(doc,caption,caption=True)


def render_docx(config,facts,narratives,figures,out,template):
    doc=Document(template)
    doc.core_properties.title=config['project_title']
    for section in doc.sections:
        page=section._sectPr.find(qn('w:pgNumType'))
        if page is not None:page.set(qn('w:start'),str(config.get('page_start',17)))
        header=section.header.paragraphs[0]
        if header.runs and '联调样章' in header.runs[-1].text:
            header.runs[-1].text='  '+config.get('header_title','输电线路路径方案报告（联调样章）')
            fonts(header.runs[-1],9)
    paragraph(doc,config.get('section_number','3.3')+'  路径方案',heading=True)
    paragraph(doc,'本次以真实规划地图中的同一起终点进行南、北两条路线联调。方案名称按地图整体上下位置确定；中部控制点用于形成不同走廊。以下图件为地图示意，现场照片待补。')
    picture(doc.add_paragraph(),figures[0]['path'],16)
    paragraph(doc,f"图 {config.get('section_number','3.3')}-1  南、北方案路径示意图",caption=True)
    by_id={f['id']:f for f in figures}
    for i,scheme in enumerate(facts['schemes']):
        paragraph(doc,f"{config.get('section_number','3.3')}.{i+1}  {scheme['name']}路径",heading=True)
        for p in narratives[scheme['id']]['paragraphs']:
            paragraph(doc,p['text'])
        images=[by_id[identifier] for identifier in scheme['figure_ids']]
        if images:figure_grid(doc,images,f"图 {config.get('section_number','3.3')}-{i+2}  {scheme['name']}沿线情况（地图示意及已确认素材）" if any(im['kind']=='approved_photo' for im in images) else f"图 {config.get('section_number','3.3')}-{i+2}  {scheme['name']}沿线情况（地图示意）")
    out=Path(out);doc.save(out)
    return str(out)


def render_pdf(docx,soffice=''):
    binary=soffice or shutil.which('soffice') or shutil.which('libreoffice')
    if not binary:return {'status':'unavailable','reason':'未安装可用的DOCX渲染器'}
    docx=Path(docx).resolve()
    with tempfile.TemporaryDirectory(prefix='rd_lo_') as profile:
        generated=Path(profile)/'rendered';generated.mkdir()
        p=subprocess.run([binary,'--headless','-env:UserInstallation='+Path(profile).as_uri(),
                          '--convert-to','pdf','--outdir',str(generated),str(docx)],
                         capture_output=True,text=True,timeout=180)
        fresh=generated/docx.with_suffix('.pdf').name
        if p.returncode or not fresh.exists():
            return {'status':'failed','reason':(p.stdout+p.stderr)[-1000:] or '本次转换未生成PDF'}
        pdf=docx.with_suffix('.pdf')
        shutil.move(str(fresh),str(pdf))
    return {'status':'completed','path':str(pdf)}


def inspect_pdf(path,folder):
    import pymupdf as fitz
    doc=fitz.open(path);folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    pages=[]
    for i,page in enumerate(doc):
        page.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(folder/f'page_{i+1:02d}.png')
        text=page.get_text()
        pages.append({'page':i+1,'text_characters':len(text),'image_count':len(page.get_image_info()),
                      'text':text})
    import re
    for old in folder.glob('page_*.png'):
        match=re.fullmatch(r'page_(\d+)\.png',old.name)
        if match and int(match.group(1))>len(doc):old.unlink()
    doc.close();return pages
