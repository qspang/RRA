"""Acceptance for the concise reference-style rich-crossing demonstration."""
import argparse,json,sys
from pathlib import Path
from docx import Document
import pymupdf
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from road_description.io import read_json,write_json,file_hash, config_root
from road_description.presentation import major_event


def main():
 p=argparse.ArgumentParser();p.add_argument('--config',required=True);p.add_argument('--baseline',default='configs/config.rich-case.resolved.json');a=p.parse_args()
 cp=Path(a.config).resolve();c=read_json(cp);out=(config_root(cp,c)/c['output_dir']).resolve()
 b=read_json(a.baseline);base=Path(b['output_dir']);checks=[]
 def yes(k,v):checks.append({'check':k,'passed':bool(v)})
 facts=read_json(out/'route_facts.json');old=read_json(base/'route_facts.json')
 yes('路线几何未为精简报告而改变',{s['id']:s['geometry_hash'] for s in facts['schemes']}=={s['id']:s['geometry_hash'] for s in old['schemes']})
 comp=read_json(out/'agents/route_comparison/result.json');r=read_json(out/'agents/route_compliance/result.json');oldr=read_json(base/'agents/route_compliance/result.json')
 yes('比较表按参考保留12类指标',len(comp['matrix']['rows'])==12)
 yes('正文未加入辅助起伏风冰成本指标',not any(x[0] in ['沿线高程起伏（辅助）','风冰区资料','模型总成本/万元'] for x in comp['matrix']['rows']))
 figures=read_json(out/'figure_manifest.json');overview=next(f for f in figures if f['id']=='overview')
 yes('总览图包含真实范围高亮及标注',overview['highlighted_area_count']>0 and len(overview['area_labels'])>0)
 for s in facts['schemes']:
  byid={e['fact_id']:e for e in s['events']};selected=[f for f in figures if f.get('scheme_id')==s['id']]
  yes(s['id']+'沿线代表图不超过6张',len(selected)<=6)
  yes(s['id']+'沿线代表图均属主要对象',all(all(major_event(byid[i]) for i in f.get('fact_ids',[]) if i in byid) for f in selected))
 summary=[]
 for sid,s in r['schemes'].items():
  old_events=oldr['schemes'][sid]['crossings'];events=s['crossings'];shown=[e for e in events if e['presentation_included']]
  yes(sid+'原始事件完整保留',[e['event_id'] for e in events]==[e['event_id'] for e in old_events])
  yes(sid+'普通道路未列入主要对象',all(major_event(e)==e['presentation_included'] for e in events))
  mapped={f.get('event_id') for f in s['figures']}
  yes(sid+'每个正文主要交点均有场景图',all(e['event_id'] in mapped for e in shown))
  yes(sid+'普通道路未逐条配图',all(e['event_id'] not in mapped for e in events if not e['presentation_included']))
  doc=Document(s['document']['docx']);text='\n'.join(p.text for p in doc.paragraphs)
  yes(sid+'没有技术附录或重复待办表','附：技术核查记录' not in text and not any('资料与复核待办' in p.text for p in doc.paragraphs))
  for category in ['高速','铁路','河流','输电']:yes(sid+'保留实际'+category+'正例',any(category in e['category'] for e in shown))
  positive=[e for e in s['ecology'] if e['inside_length_m']>.05]
  yes(sid+'保留生态实际相交',bool(positive))
  if any(e['angle_status']=='conflict' for e in events):yes(sid+'真实规则冲突未被删去','尚未满足当前规则' in text)
  summary.append({'scheme':s['name'],'all_events':len(events),'main_events':len(shown),'sidecar_events':len(events)-len(shown),'scene_images':len({f['path'] for f in s['figures']})})
 assembly=read_json(out/'agents/assembly/result.json')['documents'][0]
 with pymupdf.open(assembly['pdf']['path']) as pdf:
  text='\n'.join(p.get_text() for p in pdf)
  for title in ['3.3.1','3.3.2','3.4.8.1','3.4.8.2','3.4.9','3.4.10','3.4.11','3.5']:yes('完整章节'+title,title in text)
  yes('本测试整册不超过30页',len(pdf)<=30);pages=len(pdf)
 q=read_json(out/'quality_review.json')
 yes('使用更新后的参考PDF',q['reference_pdf_sha256']==file_hash(c['reference_pdf']) and q['reference_style']['reference_page_count']==27)
 yes('交付完整性检查通过',read_json(out/'delivery_manifest.json')['status']=='passed')
 yes('全文逐页检查通过',q['status']=='passed')
 result={'status':'passed' if all(x['passed'] for x in checks) else 'failed','checks':checks,'schemes':summary,'pages':pages,'full_report':assembly}
 write_json(out.parent/'参照版验收.json',result);print(json.dumps(result,ensure_ascii=False,indent=2))
 if result['status']!='passed':raise SystemExit(1)

if __name__=='__main__':main()
